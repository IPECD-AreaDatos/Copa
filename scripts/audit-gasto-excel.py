"""Read-only, reproducible extraction of every cell and chart in the four references.

Run with the bundled Python runtime. Never saves/changes a source workbook.
"""
import argparse
import hashlib
import json
import re
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

FILES = [
    'desagregado ministerios ene-jun.xlsx',
    'desagregado ministerios ene-jun consolida bs y serv.xlsx',
    'Ejecucion Comp - fte 10 ene a jun.xlsx',
    'desagregado partidas global ene-jun.xlsx',
]


def extract(directory):
    books = []
    for index, filename in enumerate(FILES):
        path = directory / filename
        wb = load_workbook(path, data_only=False)
        cached = load_workbook(path, data_only=True)
        sheets = []
        for ws in wb:
            values = cached[ws.title]
            cells = []
            for row in ws:
                for cell in row:
                    if cell.value is None:
                        continue
                    item = {'ref': cell.coordinate, 'row': cell.row, 'col': cell.column,
                            'value': values[cell.coordinate].value, 'format': cell.number_format}
                    if cell.data_type == 'f':
                        item['formula'] = cell.value
                    if cell.comment:
                        item['comment'] = cell.comment.text
                    if cell.font.bold:
                        item['bold'] = True
                    if cell.fill.fgColor.type == 'rgb' and cell.fill.fgColor.rgb not in ('00000000', 'FFFFFFFF'):
                        item['fill'] = '#' + cell.fill.fgColor.rgb[-6:]
                    cells.append(item)
            charts = []
            for chart in ws._charts:
                xml = chart.title.to_tree() if chart.title else chart.to_tree()
                title = ' '.join(n.text for n in xml.iter() if n.tag.split('}')[-1] == 't' and n.text) or f'Gráfico {len(charts)+1}'
                series = []
                for s in chart.series:
                    def ref_of(obj):
                        if obj is None:
                            return None
                        return getattr(getattr(obj, 'numRef', None), 'f', None) or getattr(getattr(obj, 'strRef', None), 'f', None)
                    def read_ref(ref):
                        if not ref or '!' not in ref:
                            return []
                        sh, rng = ref.rsplit('!', 1)
                        sh = sh.strip("'").replace("''", "'")
                        target = cached[sh][rng]
                        if hasattr(target, 'value'):
                            return [target.value]
                        return [c.value for row in target for c in row]
                    cat_ref, val_ref = ref_of(s.cat), ref_of(s.val)
                    tx_ref = getattr(getattr(s.tx, 'strRef', None), 'f', None)
                    name = (read_ref(tx_ref) or [getattr(s.tx, 'v', None) or 'Monto'])[0]
                    series.append({'name': str(name), 'categories': read_ref(cat_ref), 'values': read_ref(val_ref),
                                   'categoryRef': cat_ref, 'valueRef': val_ref})
                charts.append({'title': title, 'type': chart.__class__.__name__, 'series': series})
            sheets.append({'id': f'{index+1}-{len(sheets)+1}', 'name': ws.title, 'range': ws.calculate_dimension(),
                           'state': ws.sheet_state, 'columns': [get_column_letter(i) for i in range(1, ws.max_column+1)],
                           'hiddenColumns': [k for k, v in ws.column_dimensions.items() if v.hidden],
                           'hiddenRows': [k for k, v in ws.row_dimensions.items() if v.hidden],
                           'cells': cells, 'charts': charts,
                           'formulaCount': sum('formula' in c for c in cells)})
        books.append({'id': str(index+1), 'name': filename, 'sha256': hashlib.sha256(path.read_bytes()).hexdigest(), 'sheets': sheets})
    return {'extractedOn': '2026-09-04', 'period': 'enero-junio', 'yearConfirmed': False, 'books': books}


def outline(data):
    for book in data['books']:
        print('\nFILE', book['name'])
        for sheet in book['sheets']:
            print('\nSHEET', sheet['id'], sheet['name'], sheet['range'], 'cells', len(sheet['cells']), 'formulas', sheet['formulaCount'], 'charts', len(sheet['charts']))
            by_row = {}
            for cell in sheet['cells']:
                by_row.setdefault(cell['row'], []).append(cell)
            prev_denom = None
            for row, cells in by_row.items():
                denoms = [re.findall(r'/SUM\(([^)]+)\)|/\+?\$([A-Z]+\$\d+)', c.get('formula', ''), re.I) for c in cells]
                denom = str([x for x in denoms if x])
                labeled = any(isinstance(c['value'], str) and c['col'] == 1 and not c['value'].strip().isdigit() for c in cells)
                if book['id'] == '3' or labeled or (denom != '[]' and denom != prev_denom):
                    print(' | '.join(f"{c['ref']}:{str(c.get('formula', c['value']))}" for c in cells))
                if denom != '[]':
                    prev_denom = denom
            for chart in sheet['charts']:
                print('CHART', chart['title'], [(s['name'], s['categoryRef'], s['valueRef']) for s in chart['series']])


if __name__ == '__main__':
    parser = argparse.ArgumentParser()
    parser.add_argument('--directory', type=Path, default=Path('/Volumes/SD/Descargas'))
    parser.add_argument('--output', type=Path)
    parser.add_argument('--outline', action='store_true')
    args = parser.parse_args()
    data = extract(args.directory)
    if args.output:
        args.output.parent.mkdir(parents=True, exist_ok=True)
        args.output.write_text(json.dumps(data, ensure_ascii=False, separators=(',', ':')), encoding='utf-8')
    if args.outline:
        outline(data)
    else:
        print(json.dumps({b['name']: {'sheets': len(b['sheets']), 'cells': sum(len(s['cells']) for s in b['sheets']), 'charts': sum(len(s['charts']) for s in b['sheets'])} for b in data['books']}, indent=2))
